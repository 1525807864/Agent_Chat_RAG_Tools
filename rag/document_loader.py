# -*- coding: utf-8 -*-
"""多格式文档加载器：把本地文件切成 LangChain Document 列表，供向量库入库。"""
from __future__ import annotations  # 允许类型注解里用 list[Document] 等现代写法
from pathlib import Path  # 跨平台路径对象
from langchain_core.documents import Document  # 统一文档结构：page_content + metadata
class DocumentLoader:
    """统一文档加载器。
    支持：PDF / TXT / MD / HTML / CSV / JSON / DOCX / XLSX
    （老式 .doc / .xls 第一版不支持，上传时会被后端拒绝。）
    """
    # 后缀 → 内部加载器名字（仅作文档说明；真正分发在 load_file 里）
    SUPPORTED_SUFFIXES = {
        ".pdf": "pypdf",
        ".txt": "text",
        ".md": "text",
        ".html": "html",
        ".csv": "csv",
        ".json": "json",
        ".docx": "docx",  # Word（新格式）
        ".xlsx": "xlsx",  # Excel（新格式）
    }
    def __init__(self, chunk_size: int = 500, chunk_overlap: int = 50):
        # 切块大小：越大上下文越多，但检索粒度变粗
        self.chunk_size = chunk_size
        # 相邻块重叠：减少「关键句刚好被切断」的概率
        self.chunk_overlap = chunk_overlap
        # 本轮load_directory()里读失败的文件说明，成功时应该是空列表
        self.load_errors: list[str] = []
    def load_file(self, file_path: str | Path) -> list[Document]:
        """加载单个文件并切块，返回 Document 列表。"""
        path = Path(file_path)  # 统一成 Path
        if not path.exists():  # 文件必须存在
            raise FileNotFoundError(f"文件不存在: {file_path}")
        suffix = path.suffix.lower()  # 小写后缀，如 .PDF → .pdf
        if suffix not in self.SUPPORTED_SUFFIXES:  # 白名单校验
            raise ValueError(
                f"不支持的文件格式: {suffix}, 支持: {list(self.SUPPORTED_SUFFIXES)}"
            )
        # 后缀 → 实际加载函数
        loaders = {
            ".pdf": self._load_pdf,
            ".txt": self._load_text,
            ".md": self._load_text,
            ".html": self._load_html,
            ".csv": self._load_csv,
            ".json": self._load_json,
            ".docx": self._load_docx,
            ".xlsx": self._load_xlsx,
        }
        return loaders[suffix](path)  # 调用对应加载器
    def load_directory(self, dir_path: str | Path, recursive: bool = True) -> list[Document]:
        """扫描目录内所有支持格式的文件并加载。"""
        path = Path(dir_path) #统一成path
        if not path.is_dir(): #目录必须存在
            raise NotADirectoryError(f"目录不存在: {dir_path}")
        self.load_errors = []#每次扫描先清空，避免上次的失败还留着
        docs: list[Document] = []  # 汇总结果
        pattern = "**/*" if recursive else "*"  # 是否递归子目录
        for file_path in path.glob(pattern):
            if not file_path.is_file(): #跳过子目录本身
                continue 
            if file_path.suffix.lower() not in self.SUPPORTED_SUFFIXES: #不支持的后缀不管
                continue 
            try:
                docs.extend(self.load_file(file_path)) #成功就加入切块
            except Exception as e:
                self.load_errors.append(f"加载 {file_path} 失败: {e}") 
        
        return docs #调用方还要检查load_errors
            
    # ---- 各格式加载器 ----
    def _load_pdf(self, path: Path) -> list[Document]:
        """PDF：用 PyPDFLoader 读页并按 chunk 切分。"""
        from langchain_community.document_loaders import PyPDFLoader
        from langchain_text_splitters import RecursiveCharacterTextSplitter
        pages = PyPDFLoader(str(path)).load()
        splitter = RecursiveCharacterTextSplitter(
            chunk_size = self.chunk_size,
            chunk_overlap = self.chunk_overlap,
        )
        return splitter.split_documents(pages)
    def _load_text(self, path: Path) -> list[Document]:
        """纯文本 / Markdown：整文件读入再按分隔符递归切分。"""
        from langchain_text_splitters import RecursiveCharacterTextSplitter
        text = path.read_text(encoding="utf-8")  # 按 UTF-8 读全文
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=self.chunk_size,
            chunk_overlap=self.chunk_overlap,
            separators=["\n\n", "\n", "。", ".", " ", ""],  # 优先按段/句切
        )
        # metadata 带上源路径，便于 RAG 引用来源
        return splitter.create_documents([text], [{"source": str(path)}])
    def _load_html(self, path: Path) -> list[Document]:
        """HTML：去掉脚本样式后抽纯文本再切块。"""
        from bs4 import BeautifulSoup
        from langchain_text_splitters import RecursiveCharacterTextSplitter
        html = path.read_text(encoding="utf-8")
        soup = BeautifulSoup(html, "lxml")
        for tag in soup(["script", "style", "nav", "footer"]):
            tag.decompose()  # 删除噪音标签
        text = soup.get_text(separator="\n", strip=True)
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=self.chunk_size, chunk_overlap=self.chunk_overlap
        )
        return splitter.create_documents([text], [{"source": str(path)}])
    def _load_csv(self, path: Path) -> list[Document]:
        """CSV：每一行变成一条 Document（适合表格问答）。"""
        import csv
        docs: list[Document] = []
        with open(path, "r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for i, row in enumerate(reader):
                content = " | ".join(f"{k}: {v}" for k, v in row.items())
                docs.append(
                    Document(
                        page_content=content,
                        metadata={"source": str(path), "row": i},
                    )
                )
        return docs
    def _load_json(self, path: Path) -> list[Document]:
        """JSON：list 则每项一块；dict 则整对象一块。"""
        import json
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, list):
            return [
                Document(
                    page_content=json.dumps(item, ensure_ascii=False),
                    metadata={"source": str(path), "index": i},
                )
                for i, item in enumerate(data)
            ]
        if isinstance(data, dict):
            return [
                Document(
                    page_content=json.dumps(data, ensure_ascii=False),
                    metadata={"source": str(path)},
                )
            ]
        return []
    def _load_docx(self, path: Path) -> list[Document]:
        """Word .docx：抽取所有非空段落，拼成文本后再切块。"""
        from docx import Document as DocxDocument  # python-docx 的 Document
        from langchain_text_splitters import RecursiveCharacterTextSplitter
        docx = DocxDocument(str(path))  # 打开 docx
        # 只保留有文字的段落，用换行拼成一篇长文
        paragraphs = [p.text.strip() for p in docx.paragraphs if p.text and p.text.strip()]
        text = "\n".join(paragraphs)
        if not text.strip():  # 空文档：返回空列表，不报错
            return []
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=self.chunk_size,
            chunk_overlap=self.chunk_overlap,
            separators=["\n\n", "\n", "。", ".", " ", ""],
        )
        return splitter.create_documents(
            [text],
            [{"source": str(path), "filetype": "docx"}],
        )
    def _load_xlsx(self, path: Path) -> list[Document]:
        """Excel .xlsx：每个工作表按「表头+行」转成文本块。
        策略：每个 sheet 先拼成一大段文字，再统一切块；
        比「一行一块」更省向量条数，也适合概括性提问。
        """
        from openpyxl import load_workbook
        from langchain_text_splitters import RecursiveCharacterTextSplitter
        wb = load_workbook(str(path), data_only=True, read_only=True)  # 只读更快
        parts: list[str] = []  # 收集各 sheet 文本
        for sheet in wb.worksheets:
            rows = list(sheet.iter_rows(values_only=True))  # 全部单元格值
            if not rows:
                continue
            # 第一行当表头（若存在）
            header = [str(c) if c is not None else "" for c in rows[0]]
            parts.append(f"[Sheet] {sheet.title}")  # 标明 sheet 名
            for r_i, row in enumerate(rows[1:], start=1):  # 从第 2 行起是数据
                cells = []
                for c_i, cell in enumerate(row):
                    # 有表头则用「列名: 值」，否则用「cN: 值」
                    key = header[c_i] if c_i < len(header) and header[c_i] else f"c{c_i}"
                    val = "" if cell is None else str(cell)
                    if val.strip():
                        cells.append(f"{key}: {val}")
                if cells:
                    parts.append(f"row{r_i}: " + " | ".join(cells))
        text = "\n".join(parts)
        if not text.strip():
            return []
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=self.chunk_size,
            chunk_overlap=self.chunk_overlap,
            separators=["\n\n", "\n", "。", ".", " ", ""],
        )
        return splitter.create_documents(
            [text],
            [{"source": str(path), "filetype": "xlsx"}],
        )