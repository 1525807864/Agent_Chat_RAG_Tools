"""把一次 UI 运行导出为 xlsx / docx / pdf。数据只来自已落盘的 steps。"""

from __future__ import annotations

import io
from typing import Any

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font

from domain.qa import QaRun


def _steps(run: QaRun) -> list[dict[str, Any]]:
    raw = (run.artifacts or {}).get("steps") or []
    return [s for s in raw if isinstance(s, dict)]


def _meta(run: QaRun) -> dict[str, str]:
    art = run.artifacts or {}
    return {
        "名称": str(art.get("name") or ""),
        "目标": str(art.get("goal") or ""),
        "地址": str(art.get("url") or ""),
        "结论": run.verdict,
        "摘要": run.summary,
        "时间": run.created_at,
    }


def to_xlsx(run: QaRun) -> bytes:
    wb = Workbook()
    cover = wb.active
    cover.title = "结论"
    cover["A1"] = "字段"
    cover["B1"] = "内容"
    cover["A1"].font = Font(bold=True)
    for i, (k, v) in enumerate(_meta(run).items(), start=2):
        cover.cell(i, 1, k)
        cover.cell(i, 2, v)
    ws = wb.create_sheet("步骤")
    headers = ["序号", "动作", "说明", "结果", "失败信息"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for s in _steps(run):
        ws.append(
            [
                s.get("i"),
                s.get("action"),
                s.get("detail"),
                "通过" if s.get("ok") else "失败",
                s.get("error") or "",
            ]
        )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_docx(run: QaRun) -> bytes:
    doc = Document()
    doc.add_heading("UI 测试执行记录", level=1)
    for k, v in _meta(run).items():
        doc.add_paragraph(f"{k}：{v}")
    doc.add_heading("步骤", level=2)
    table = doc.add_table(rows=1, cols=5)
    for i, h in enumerate(["序号", "动作", "说明", "结果", "失败信息"]):
        table.rows[0].cells[i].text = h
    for s in _steps(run):
        row = table.add_row().cells
        row[0].text = str(s.get("i") or "")
        row[1].text = str(s.get("action") or "")
        row[2].text = str(s.get("detail") or "")
        row[3].text = "通过" if s.get("ok") else "失败"
        row[4].text = str(s.get("error") or "")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_pdf(run: QaRun) -> bytes:
    """PDF 用 reportlab；未安装时抛出明确错误。"""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.pdfgen import canvas
    except ImportError as exc:
        raise RuntimeError("导出 PDF 需要安装 reportlab：pip install reportlab") from exc

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setFont("STSong-Light", 12)
    y = 800
    c.drawString(40, y, "UI 测试执行记录")
    y -= 24
    for k, v in _meta(run).items():
        c.drawString(40, y, f"{k}：{v}"[:90])
        y -= 18
        if y < 80:
            c.showPage()
            c.setFont("STSong-Light", 12)
            y = 800
    y -= 10
    c.drawString(40, y, "步骤")
    y -= 18
    for s in _steps(run):
        line = (
            f"{s.get('i')}. {s.get('action')} | "
            f"{'通过' if s.get('ok') else '失败'} | {s.get('detail')}"
        )
        c.drawString(40, y, line[:90])
        y -= 16
        if y < 60:
            c.showPage()
            c.setFont("STSong-Light", 12)
            y = 800
    c.save()
    return buf.getvalue()


def export_bytes(run: QaRun, fmt: str) -> tuple[bytes, str, str]:
    """返回 (字节, mime, 文件名)。"""
    fmt = (fmt or "xlsx").lower()
    if fmt == "xlsx":
        return (
            to_xlsx(run),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"{run.id}.xlsx",
        )
    if fmt == "docx":
        return (
            to_docx(run),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            f"{run.id}.docx",
        )
    if fmt == "pdf":
        return to_pdf(run), "application/pdf", f"{run.id}.pdf"
    raise ValueError("fmt 只能是 xlsx、docx、pdf")
